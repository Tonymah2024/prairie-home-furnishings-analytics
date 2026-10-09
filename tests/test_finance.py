"""Financial statement reconciliation: double entry, trial balances, A = L + E, retained earnings, cash, inventory,
revenue/COGS, deposits, forecast integrity under changed assumptions, Excel model, labelling."""
from dataclasses import replace

import numpy as np
import openpyxl
import pandas as pd
import pytest

from src.config import FIN_OUT, OUT, TABLES
from src.finance import forecast as F
from src.finance import ledger as LG
from src.finance import policy as P
from src.finance import statements as S
from src.finance import validation as V

TOL = 0.01


@pytest.fixture(scope="module")
def fin():
    import json
    K = json.loads((TABLES / "findings.json").read_text())["metrics"]
    L = LG.build()
    St = S.build(L)
    return L, St, F.build(L, St, K)


def test_every_journal_entry_balances(fin):
    gl = fin[0]["gl"]
    je = gl.groupby("je_id")[["debit", "credit"]].sum()
    assert ((je.debit - je.credit).abs() < 0.005).all()


def test_unbalanced_entry_rejected():
    with pytest.raises(ValueError):
        LG.Book().je("X", "2024-01-01", "Customer payments", "X", "bad", [("1000", 10.0, None), ("2350", -9.0, None)])


def test_trial_balances_and_subledgers(fin):
    St = fin[1]
    assert (St["tb"].groupby("month").apply(lambda x: abs(x.debit.sum() - x.credit.sum()), include_groups=False) < TOL).all()
    assert St["subledger"].difference.abs().max() < TOL


def test_accounting_equation(fin):
    St, FC = fin[1], fin[2]
    assert St["bs_m"].loc["check"].abs().max() < TOL and abs(St["bs_open"]["check"]) < TOL
    assert FC["summary"].max_abs_bs_check.max() < TOL


def test_retained_earnings_movements(fin):
    St = fin[1]
    bs, is_y, cf_y = St["bs_m"], St["is_y"], St["cf_y"]
    assert abs(bs.at["retained", "2024-12"] - (St["bs_open"]["retained"] + is_y.at["net_income", "2024"] + cf_y.at["dividends", "2024"])) < TOL
    assert abs(bs.at["retained", "2025-12"] - (bs.at["retained", "2024-12"] + is_y.at["net_income", "2025"] + cf_y.at["dividends", "2025"])) < TOL
    assert abs(bs.at["cpe", "2026-09"] - is_y.at["net_income", "2026"]) < TOL


def test_cash_flow_reconciles(fin):
    St = fin[1]
    for k in ("cf_m", "cf_q", "cf_y"):
        assert St[k].loc["check"].abs().max() < TOL
    assert np.allclose(St["direct"].sum().values, St["cf_m"].loc["net_change"].values, atol=TOL)
    assert V.daily_cash(fin[0]["gl"]).min() >= 0


def test_inventory_and_revenue_reconcile(fin):
    L, St = fin[0], fin[1]
    iv = St["inv_val"]
    assert np.allclose(iv.stock_value_from_movements, St["bs_m"].loc["inv_stock"], atol=TOL)
    assert np.allclose(iv.allocated_value_from_movements, St["bs_m"].loc["inv_alloc"], atol=TOL)
    assert V.revenue_cogs_reconciliation(L, St).difference.abs().max() < TOL
    assert abs(V.bookings_to_revenue_bridge(L, St).cad.iloc[-1]) < TOL


def test_revenue_recognised_at_handoff_not_order(fin):
    L = fin[0]
    gl = L["gl"]
    sales = gl[gl.account.isin(P.REV_ACCT.values()) & (gl.source == "Order completion")]
    o = L["T"]["orders"].set_index("order_id")
    assert (sales.date.values == o.loc[sales.reference, "fulfilled_date"].values).all()
    assert not gl[gl.account.isin(P.REV_ACCT.values()) & (gl.source != "Order completion") & ~gl.is_closing].shape[0]


@pytest.mark.parametrize("seed", range(4))
def test_forecast_integrated_under_random_assumptions(fin, seed):
    rng = np.random.default_rng(seed)
    sc = F.Scenario(volume_change=rng.uniform(-.4, .4), margin_pts=rng.uniform(-.06, .06), slow_reduction=rng.uniform(0, .6), delay_reduction=rng.uniform(-.6, .6),
                    consult_conv_change=rng.uniform(-.4, .4), dpo_change_days=rng.uniform(-25, 25), opex_change=rng.uniform(-.25, .25),
                    discount_change=rng.uniform(-.04, .06), rate_change=rng.uniform(-.03, .05))
    r = F.run(sc, fin[2]["base"])
    assert r.loc["check_bs"].iloc[1:].abs().max() < TOL and r.loc["check_cash"].iloc[1:].abs().max() < TOL


def test_scenario_directions(fin):
    b = fin[2]["base"]
    s0 = F.summary(F.run(F.SCENARIOS["Base"], b))
    s = lambda **k: F.summary(F.run(replace(F.SCENARIOS["Base"], **k), b))
    assert s(volume_change=.05)["net_income"] > s0["net_income"]
    assert s(margin_pts=.02)["gross_profit"] > s0["gross_profit"]
    assert s(dpo_change_days=10)["closing_cash"] > s0["closing_cash"]
    assert s(opex_change=.05)["net_income"] < s0["net_income"]
    assert s(slow_reduction=.1)["closing_cash"] > s0["closing_cash"]
    sm = fin[2]["summary"]
    assert sm.at["Conservative", "net_income"] < sm.at["Base", "net_income"] < sm.at["Optimistic", "net_income"]


def test_forecast_opens_from_ledger(fin):
    f, bs = fin[2]["results"]["Base"], fin[1]["bs_m"]
    for k in ("cash", "inv_stock", "inv_alloc", "deposits", "retained", "cpe"):
        assert abs(f.iloc[:, 0][k] - bs.at[k, "2026-09"]) < TOL, k


def test_excel_model_formulas_and_cached_values(fin):
    p = FIN_OUT / "Prairie_Home_Integrated_Financial_Model.xlsx"
    if not p.exists():
        pytest.skip("run the pipeline first")
    ws = openpyxl.load_workbook(p)["Forecast"]
    rows = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
    assert set(F.ROW) <= set(rows)
    assert all(str(ws.cell(row=rows["net_income"], column=c).value).startswith("=") for c in range(4, 16))
    v = openpyxl.load_workbook(p, data_only=True)["Forecast"]
    cached = [v.cell(row=rows["net_income"], column=c).value for c in range(4, 16)]
    if all(x is not None for x in cached):
        assert np.allclose(cached, fin[2]["results"]["Base"].loc["net_income"].iloc[1:].values, atol=TOL)


def test_validation_checks_and_labels():
    ch = pd.read_csv(TABLES / "fin_validation_checks.csv")
    assert (ch.result == "PASS").all(), ch
    ok, detail = V.label_audit(OUT)
    assert ok, detail
