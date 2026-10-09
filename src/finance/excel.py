"""Phase 18 - Excel deliverables. The integrated model is formula-driven: historical subtotals, the forecast,
scenario selection and all checks are live Excel formulas, generated from the same templates the Python engine
evaluates (forecast.ROWS). All workbooks carry the SIMULATED-data banner on every sheet."""
from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from src.config import BRAND
from src.finance import forecast as F
from src.finance import policy as P
from src.finance import ratios as R
from src.finance import statements as S

NAVY = BRAND["navy"].lstrip("#")
BLUE = BRAND["brand_dark"].lstrip("#")
FMT = '#,##0;(#,##0);"-"'
FMT2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
F_IN = Font(color="0000FF")
FILL_IN = PatternFill("solid", fgColor="FFF8DC")
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_SUB = PatternFill("solid", fgColor="EAF2FC")
FILL_BAN = PatternFill("solid", fgColor="FFF4E5")
THIN = Border(top=Side(style="thin", color="999999"))
BANNER = "⚠ " + P.FIN_LABEL + ". Synthetic data generated for a VisiVault Analytics proof of concept (seed 20240102). Fictional retailer."


def banner(ws, title, sub=None, width=16):
    ws["A1"] = BANNER
    ws["A1"].font = Font(bold=True, color="7A4A00")
    ws["A1"].fill = FILL_BAN
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=14, color=BLUE)
    if sub:
        ws["A3"] = sub
        ws["A3"].font = Font(italic=True, color="52514E")
    ws.sheet_view.showGridLines = False


def header(ws, row, values, start_col=1):
    for j, v in enumerate(values):
        c = ws.cell(row=row, column=start_col + j, value=v)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = FILL_HDR
        c.alignment = Alignment(horizontal="center" if j else "left", wrap_text=True)


def style_sub(ws, row, ncols, bold=True):
    for j in range(1, ncols + 1):
        c = ws.cell(row=row, column=j)
        c.font = Font(bold=bold)
        c.fill = FILL_SUB
        c.border = THIN


def write_df(ws, df: pd.DataFrame, r0=5, index=True, fmt=FMT, subtotal_rows=(), widths=None):
    cols = ([df.index.name or ""] if index else []) + [str(c) for c in df.columns]
    header(ws, r0, cols)
    for i, (idx, row) in enumerate(df.iterrows()):
        r = r0 + 1 + i
        vals = ([idx] if index else []) + list(row.values)
        for j, v in enumerate(vals):
            if isinstance(v, float) and pd.isna(v):
                v = None
            c = ws.cell(row=r, column=j + 1, value=v)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                c.number_format = fmt
        if idx in subtotal_rows:
            style_sub(ws, r, len(vals))
    ws.freeze_panes = ws.cell(row=r0 + 1, column=2 if index else 1)
    for j in range(1, len(cols) + 1):
        ws.column_dimensions[L(j)].width = (widths or {}).get(j, 44 if j == 1 else 13)
    return r0 + len(df) + 1


# ======================================================================================== integrated model
IS_ORDER = [k for k, *_ in S.IS_LINES]
IS_SUB, BS_SUB, CF_SUB = S.IS_SUB, S.BS_SUB, S.CF_SUB

FC_SECTIONS = [
    ("INCOME STATEMENT (forecast)", ["gross_sales", "discounts", "returns", "net_sales", "other_rev", "net_rev", "cogs_merch", "freight_in", "inv_var", "cogs",
                                     "gross_profit", "emp_comp", "occupancy", "marketing", "g_and_a", "warehouse", "fees", "opex", "ebitda", "da", "ebit",
                                     "interest", "pbt", "tax", "net_income"]),
    ("BALANCE SHEET (forecast; column C = actual 30-Sep-2026 from the general ledger)",
     ["cash", "ar", "inv_stock", "inv_alloc", "prepaid", "gst_rec", "current_assets", "ppe_gross", "acc_dep", "ppe_net", "total_assets", "ap", "accrued", "deposits",
      "agent", "gst_pay", "tax_payable", "revolver", "term_current", "current_liabilities", "term_lt", "total_liabilities", "share_capital", "retained", "cpe",
      "total_equity", "total_le"]),
    ("CASH FLOW STATEMENT (indirect method)", ["cf_ni", "cf_da", "d_ar", "d_inv", "d_ap", "d_deposits", "d_other", "cfo", "cfi", "term_repay", "rev_draw",
                                               "rev_repay", "dividends", "cff", "net_change"]),
    ("CHECKS", ["check_bs", "check_cash"]),
    ("SUPPORTING SCHEDULES - bookings, recognition lags & order backlog", ["bookings", "w1e", "w2e", "w3e", "w4e", "cancel_eff", "backlog"]),
    ("SUPPORTING SCHEDULES - payroll, inventory investment & capex", ["payroll", "stock_target", "purchases", "capex_cum"]),
    ("SUPPORTING SCHEDULES - working capital & financing", ["gst_net", "gst_remit", "working_capital", "pre_cash", "term_loan"]),
    ("SUPPORTING SCHEDULES - tax & dividends", ["ytd_pbt", "lcf", "tax_ytd", "prior_tax", "tax_paid", "ytd_ni"]),
]
SUBTOTAL_KEYS = set(IS_SUB) | set(BS_SUB) | {"cfo", "cff", "net_change", "net_sales", "net_rev", "gross_profit", "net_income", "check_bs", "check_cash"}
FLOW_KEYS = {k for _, ks in FC_SECTIONS[:1] + FC_SECTIONS[2:3] for k in ks} | {"bookings", "payroll", "purchases", "gst_remit", "tax_paid"}


class Model:
    def __init__(self, L_, S_, FC):
        self.L, self.S, self.FC = L_, S_, FC
        self.wb = openpyxl.Workbook()
        self.pos = {}          # key -> (sheet, row)
        self.apos = {}         # assumption -> absolute ref

    # ------------------------------------------------------------------ historical sheets
    def hist_balances(self):
        ws = self.wb.create_sheet("Hist_Balances")
        banner(ws, "Month-end account balances from the general ledger (debit +, credit -)", "Post-closing at year ends. Each column must sum to zero (trial balance).", 36)
        bal = self.S["bal"]
        months = [str(m) for m in bal.index]
        header(ws, 5, ["Account", "Name"] + months)
        self.bal_col = {m: L(3 + i) for i, m in enumerate(months)}
        self.bal_row = {}
        for i, a in enumerate(P.ACCOUNTS):
            r = 6 + i
            self.bal_row[a] = r
            ws.cell(row=r, column=1, value=a)
            ws.cell(row=r, column=2, value=P.ACCOUNTS[a]["name"])
            for j, m in enumerate(months):
                ws.cell(row=r, column=3 + j, value=float(bal.at[pd.Period(m, "M"), a])).number_format = FMT2
        r = 6 + len(P.ACCOUNTS)
        ws.cell(row=r, column=2, value="Check: sum of balances (must be 0)")
        for j in range(len(months)):
            c = L(3 + j)
            ws.cell(row=r, column=3 + j, value=f"=ROUND(SUM({c}6:{c}{r - 1}),2)").number_format = FMT2
        style_sub(ws, r, 2 + len(months))
        self.tb_check = f"Hist_Balances!$C${r}:${L(2 + len(months))}${r}"
        ws.column_dimensions["A"].width = 9
        ws.column_dimensions["B"].width = 46
        ws.freeze_panes = "C6"

    def hist_statement(self, name, title, df, order, subs, labels, extra=None):
        ws = self.wb.create_sheet(name)
        banner(ws, title, "Monthly values from the general ledger; subtotal rows are live formulas. FY columns sum the months.", 40)
        months = list(df.columns)
        years = sorted({m[:4] for m in months})
        header(ws, 5, ["Line", "Key"] + months + [f"FY{y}" if y != "2026" else "YTD Sep-2026" for y in years])
        rows = {k: 6 + i for i, k in enumerate(order)}
        flows = name != "Hist_BS"
        for k in order:
            r = rows[k]
            ws.cell(row=r, column=1, value=labels[k])
            ws.cell(row=r, column=2, value=k).font = Font(color="898781", size=8)
            for j, m in enumerate(months):
                col = L(3 + j)
                if k in subs:
                    v = "=" + "+".join(f"{col}{rows[x]}" for x in subs[k])
                elif extra and k in extra:
                    v = extra[k](j, col, rows)
                else:
                    v = float(df.at[k, m])
                ws.cell(row=r, column=3 + j, value=v).number_format = FMT
            for y_i, y in enumerate(years):
                cols = [L(3 + j) for j, m in enumerate(months) if m.startswith(y)]
                c = ws.cell(row=r, column=3 + len(months) + y_i)
                if flows and k not in ("open_cash", "close_cash", "bs_cash", "check"):
                    c.value = f"=SUM({cols[0]}{r}:{cols[-1]}{r})"
                elif k == "open_cash":
                    c.value = f"={cols[0]}{r}"
                else:
                    c.value = f"={cols[-1]}{r}"
                c.number_format = FMT
            if k in subs or k in ("check", "close_cash"):
                style_sub(ws, r, 2 + len(months) + len(years))
        ws.column_dimensions["A"].width = 52
        ws.column_dimensions["B"].width = 11
        for j in range(len(months) + len(years)):
            ws.column_dimensions[L(3 + j)].width = 11.5
        ws.freeze_panes = "C6"
        return ws, rows, {m: L(3 + j) for j, m in enumerate(months)}

    def historicals(self):
        S_ = self.S
        self.hist_balances()
        _, self.is_rows, self.is_cols = self.hist_statement("Hist_IS", "Historical income statements - illustrative Calgary management accounts (CAD)",
                                                            S_["is_m"], IS_ORDER, IS_SUB, S.IS_LABEL)
        bs_order = [k for k, *_ in S.BS_LINES]
        chk = {"check": lambda j, col, rows: f"=ROUND({col}{rows['total_assets']}-{col}{rows['total_le']},2)"}
        _, self.bs_rows, self.bs_cols = self.hist_statement("Hist_BS", "Historical balance sheets - month end (CAD)", S_["bs_m"], bs_order, BS_SUB, S.BS_LABEL, chk)
        cf_order = [k for k, *_ in S.CF_LINES]
        bs_rows, bs_cols = self.bs_rows, self.bs_cols
        months = list(S_["cf_m"].columns)
        open0 = float(S_["bs_open"]["cash"])
        ex = {
            "open_cash": lambda j, col, rows: open0 if j == 0 else f"={L(2 + j)}{rows['close_cash']}",
            "close_cash": lambda j, col, rows: f"={col}{rows['open_cash']}+{col}{rows['net_change']}",
            "bs_cash": lambda j, col, rows: f"=Hist_BS!{bs_cols[months[j]]}{bs_rows['cash']}",
            "check": lambda j, col, rows: f"=ROUND({col}{rows['close_cash']}-{col}{rows['bs_cash']},2)",
        }
        self.hist_statement("Hist_CF", "Historical cash flow statements - indirect method (CAD)", S_["cf_m"], cf_order, CF_SUB, S.CF_LABEL, ex)

    # ------------------------------------------------------------------ assumptions & drivers
    def assumptions(self):
        ws = self.wb.create_sheet("Assumptions", 1)
        banner(ws, "Forecast assumptions - choose a scenario, or edit the Custom column", "Blue cells on yellow are inputs. Column H (Active) feeds every forecast formula.", 9)
        ws["B4"] = "Active scenario  >>"
        ws["B4"].font = Font(bold=True)
        ws["C4"] = "Base"
        ws["C4"].font = Font(bold=True, color="0000FF", size=12)
        ws["C4"].fill = FILL_IN
        dv = DataValidation(type="list", formula1='"Conservative,Base,Optimistic,Custom"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add("C4")
        header(ws, 7, ["Key", "Scenario driver", "Unit", "Conservative", "Base", "Optimistic", "Custom", "Active"])
        r = 8
        for fld, lab, unit in F.SCENARIO_FIELDS:
            ws.cell(row=r, column=1, value=fld).font = Font(color="898781", size=8)
            ws.cell(row=r, column=2, value=lab)
            ws.cell(row=r, column=3, value=unit)
            for j, name in enumerate(["Conservative", "Base", "Optimistic", "Base"]):
                c = ws.cell(row=r, column=4 + j, value=float(getattr(F.SCENARIOS[name], fld)))
                c.font, c.fill, c.number_format = F_IN, FILL_IN, "0.000"
            ws.cell(row=r, column=8, value=f"=INDEX(D{r}:G{r},MATCH($C$4,$D$7:$G$7,0))").number_format = "0.000"
            self.apos[fld] = f"Assumptions!$H${r}"
            r += 1
        r += 1
        ws.cell(row=r, column=2, value="Scenario notes").font = Font(bold=True)
        for name, note in F.SCENARIO_NOTES.items():
            r += 1
            ws.cell(row=r, column=2, value=name).font = Font(bold=True)
            ws.cell(row=r, column=3, value=note)
        r += 2
        header(ws, r, ["Key", "Calibrated / policy parameter (TTM Oct-2025 to Sep-2026 unless stated)", "", "", "", "", "", "Value"])
        for k, v in self.FC["base"]["A"].items():
            r += 1
            ws.cell(row=r, column=1, value=k).font = Font(color="898781", size=8)
            ws.cell(row=r, column=2, value=F.FIXED_LABELS.get(k, k))
            c = ws.cell(row=r, column=8, value=float(v))
            c.font, c.fill = F_IN, FILL_IN
            c.number_format = "0.0000" if abs(v) < 10 else FMT
            self.apos[k] = f"Assumptions!$H${r}"
        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 58
        for c in "CDEFG":
            ws.column_dimensions[c].width = 13
        ws.column_dimensions["H"].width = 14

    def drivers(self):
        ws = self.wb.create_sheet("Drivers", 2)
        banner(ws, "Monthly drivers (inputs) - base volumes from Model 1, staffing plan, lease, marketing plan, capex plan, inventory policy",
               "Edit any blue cell; the forecast recalculates. Column C = actual 30-Sep-2026 position (not used by drivers).", 16)
        self._month_header(ws, 5)
        Dv = self.FC["base"]["drivers"]
        labels = dict(F.DRIVER_ROWS)
        for i, k in enumerate(Dv.index):
            r = 6 + i
            self.pos[k] = ("Drivers", r)
            ws.cell(row=r, column=1, value=k).font = Font(color="898781", size=8)
            ws.cell(row=r, column=2, value=labels[k])
            for t in range(1, F.H + 1):
                c = ws.cell(row=r, column=3 + t, value=float(Dv.at[k, t]))
                c.font, c.fill = F_IN, FILL_IN
                c.number_format = "0.00" if k in ("cover_k",) else ("0" if k.startswith(("is_", "q_")) else FMT)
        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 62
        ws.freeze_panes = "D6"

    def _month_header(self, ws, row):
        header(ws, row, ["Key", "Line", "Sep-2026 actual"] + [m.strftime("%b-%Y") for m in F.FC_MONTHS] + ["12-month total"])
        for j in range(3, 17):
            ws.column_dimensions[L(j)].width = 12.5

    # ------------------------------------------------------------------ forecast
    def _ref(self, key, t, suf, home):
        sheet, row = self.pos[key]
        pre = "" if sheet == home else f"{sheet}!"
        if not suf:
            return f"{pre}{L(3 + t)}{row}"
        if suf.startswith("@"):
            return f"{pre}${L(3 + int(suf[1:]))}${row}"
        k = int(suf[1:-1])
        tt = t + k
        if tt < 0:
            return "0"                         # before the opening month: zero (same rule as the Python engine)
        if tt > F.H:
            tt -= F.H
        return f"{pre}{L(3 + tt)}{row}"

    def to_excel(self, tpl, t, home="Forecast"):
        def sub(m):
            dollar, key, suf = m.groups()
            return self.apos[key] if dollar else self._ref(key, t, suf, home)
        x = F.TOK.sub(sub, tpl)
        x = x.replace("**", "^").replace("CEIL(", "CEILING(")
        return "=" + x

    def opening_formula(self, key):
        bc = self.bal_col["2026-09"]
        b = lambda *accts: "+".join(f"Hist_Balances!${bc}${self.bal_row[a]}" for a in accts)
        sep = self.is_cols["2026-09"]
        jan = self.is_cols["2026-01"]
        m = {
            "inv_stock": f"={b('1200')}", "inv_alloc": f"={b('1210')}", "ar": f"={b('1100', '1110', '1120')}", "prepaid": f"={b('1300')}",
            "deposits": f"=-{b('2350')}", "gst_net": f"=-({b('1400', '2300')})", "ap": f"=-({b('2000', '2010')})", "accrued": f"=-{b('2100')}",
            "agent": f"=-({b('2360', '2370')})", "tax_payable": f"=-{b('2400')}", "ppe_gross": f"={b('1500', '1510', '1520')}", "acc_dep": f"={b('1590')}",
            "cash": f"={b('1000')}", "term_loan": f"=-{b('2700')}", "revolver": f"=-{b('2600')}", "share_capital": f"=-{b('3000')}",
            "retained": f"=-({b('3200', '3300')})", "cpe": f"=Hist_BS!${self.bs_cols['2026-09']}${self.bs_rows['cpe']}",
            "ytd_pbt": f"=SUM(Hist_IS!${jan}${self.is_rows['pbt']}:${sep}${self.is_rows['pbt']})",
            "tax_ytd": f"=-SUM(Hist_IS!${jan}${self.is_rows['tax']}:${sep}${self.is_rows['tax']})",
            "ytd_ni": f"=SUM(Hist_IS!${jan}${self.is_rows['net_income']}:${sep}${self.is_rows['net_income']})",
        }
        return m.get(key)

    def forecast_sheet(self):
        missing = set(F.ROW) - {k for _, ks in FC_SECTIONS for k in ks}
        assert not missing, f"forecast rows missing from the Excel layout: {missing}"
        ws = self.wb.create_sheet("Forecast", 3)
        banner(ws, "Integrated 12-month forecast Oct-2026 to Sep-2027 (FORECAST - not historical fact)",
               "Every cell in columns D:O is a formula driven by Assumptions (active scenario) and Drivers. Column C links to the historical ledger.", 16)
        self._month_header(ws, 5)
        r = 6
        layout = []
        for title, keys in FC_SECTIONS:
            r += 1
            ws.cell(row=r, column=1, value=title).font = Font(bold=True, color=BLUE, size=11)
            r += 1
            for k in keys:
                self.pos[k] = ("Forecast", r)
                layout.append((k, r))
                r += 1
        for k, r in layout:
            _, lab, _, tpl, _ = F.ROW[k]
            ws.cell(row=r, column=1, value=k).font = Font(color="898781", size=8)
            ws.cell(row=r, column=2, value=lab)
            if k in F.OPENING:
                f = self.opening_formula(k)
                v = f if f else float(self.FC["base"]["opening"][k])
                c = ws.cell(row=r, column=3, value=v)
                c.number_format = FMT
                if not f:
                    c.font, c.fill = F_IN, FILL_IN
            for t in range(1, F.H + 1):
                fmt = "0.0000" if k in ("w1e", "w2e", "w3e", "w4e", "cancel_eff") else FMT
                ws.cell(row=r, column=3 + t, value=self.to_excel(tpl, t)).number_format = fmt
            if k in FLOW_KEYS:
                ws.cell(row=r, column=16, value=f"=SUM(D{r}:O{r})").number_format = FMT
            if k in SUBTOTAL_KEYS:
                style_sub(ws, r, 16)
        # BS presentation rows that also need a t=0 value for readability
        for k in ["gst_rec", "current_assets", "ppe_net", "total_assets", "gst_pay", "term_current", "current_liabilities",
                  "term_lt", "total_liabilities", "total_equity", "total_le", "check_bs"]:
            _, row = self.pos[k]
            ws.cell(row=row, column=3, value=f"=Hist_BS!${self.bs_cols['2026-09']}${self.bs_rows.get(k, self.bs_rows['check'] if k == 'check_bs' else 0)}").number_format = FMT
        for k in ["check_bs", "check_cash"]:
            _, row = self.pos[k]
            ws.conditional_formatting.add(f"D{row}:O{row}", CellIsRule(operator="notBetween", formula=["-0.01", "0.01"], fill=PatternFill("solid", fgColor="F8D7DA")))
        ws.column_dimensions["A"].width = 13
        ws.column_dimensions["B"].width = 58
        ws.freeze_panes = "D6"
        self.fc_ws = ws

    def scenario_sheet(self):
        ws = self.wb.create_sheet("Scenario_Results", 4)
        banner(ws, "Scenario comparison and sensitivity analysis (values computed by the Python engine at build time)",
               "The Forecast sheet recalculates live for the scenario selected on Assumptions; this sheet stores all three for comparison.", 12)
        sm = self.FC["summary"].drop(columns="assumptions").T
        nr = write_df(ws, sm, 5, fmt=FMT)
        r = nr + 2
        ws.cell(row=r, column=1, value="Sensitivity analysis (one driver at a time vs Base)").font = Font(bold=True, color=BLUE)
        sens = self.FC["sensitivity"].set_index("scenario")
        write_df(ws, sens, r + 1, fmt=FMT)
        ws.column_dimensions["A"].width = 40
        # live summary of the active scenario
        r = ws.max_row + 3
        ws.cell(row=r, column=1, value="LIVE: active scenario (formulas on the Forecast sheet)").font = Font(bold=True, color=BLUE)
        live = [("Net revenue", "SUM(Forecast!D{0}:O{0})", "net_rev"), ("Net income", "SUM(Forecast!D{0}:O{0})", "net_income"),
                ("Closing cash Sep-2027", "Forecast!O{0}", "cash"), ("Minimum month-end cash", "MIN(Forecast!D{0}:O{0})", "cash"),
                ("Peak revolver", "MAX(0,MAX(Forecast!D{0}:O{0}))", "revolver"), ("Customer deposits Sep-2027", "Forecast!O{0}", "deposits"),
                ("Max |balance-sheet check|", "MAX(ABS(MIN(Forecast!D{0}:O{0})),ABS(MAX(Forecast!D{0}:O{0})))", "check_bs")]
        for i, (lab, f, k) in enumerate(live):
            ws.cell(row=r + 1 + i, column=1, value=lab)
            ws.cell(row=r + 1 + i, column=2, value="=" + f.format(self.pos[k][1])).number_format = FMT
        self.live_rows = {k: r + 1 + i for i, (_, _, k) in enumerate(live)}

    def checks_sheet(self):
        ws = self.wb.create_sheet("Checks", 5)
        banner(ws, "Model integrity checks (live formulas - every value must be 0 / TRUE)", None, 6)
        bs_last = L(2 + len(self.S["bs_m"].columns))
        cf_last = L(2 + len(self.S["cf_m"].columns))
        items = [
            ("Trial balance: max |sum of account balances| (Dec-2023 to Sep-2026)", f"=MAX(ABS(MIN({self.tb_check})),ABS(MAX({self.tb_check})))"),
            ("Historical balance sheet: max |assets - liabilities - equity|", f"=MAX(ABS(MIN(Hist_BS!C{self.bs_rows['check']}:{bs_last}{self.bs_rows['check']})),ABS(MAX(Hist_BS!C{self.bs_rows['check']}:{bs_last}{self.bs_rows['check']})))"),
            ("Historical cash flow: max |closing cash - balance-sheet cash|", "=MAX(ABS(MIN(Hist_CF!C{0}:{1}{0})),ABS(MAX(Hist_CF!C{0}:{1}{0})))".format(len(S.CF_LINES) + 5, cf_last)),
            ("Forecast balance sheet: max |assets - liabilities - equity|", "=MAX(ABS(MIN(Forecast!D{0}:O{0})),ABS(MAX(Forecast!D{0}:O{0})))".format(self.pos["check_bs"][1])),
            ("Forecast cash flow: max |closing cash - balance-sheet cash|", "=MAX(ABS(MIN(Forecast!D{0}:O{0})),ABS(MAX(Forecast!D{0}:O{0})))".format(self.pos["check_cash"][1])),
            ("Forecast opening cash = historical closing cash (Sep-2026)", f"=ROUND(Forecast!C{self.pos['cash'][1]}-Hist_BS!{self.bs_cols['2026-09']}{self.bs_rows['cash']},2)"),
        ]
        header(ws, 5, ["Check", "Value", "Pass?"])
        for i, (lab, f) in enumerate(items):
            r = 6 + i
            ws.cell(row=r, column=1, value=lab)
            ws.cell(row=r, column=2, value=f).number_format = FMT2
            ws.cell(row=r, column=3, value=f"=ABS(B{r})<0.01")
        r = 6 + len(items)
        ws.cell(row=r, column=1, value="ALL CHECKS PASS").font = Font(bold=True)
        ws.cell(row=r, column=3, value=f"=AND(C6:C{r - 1})").font = Font(bold=True)
        self.all_checks = f"Checks!C{r}"
        ws.column_dimensions["A"].width = 72
        ws.column_dimensions["B"].width = 14

    def readme(self):
        ws = self.wb.active
        ws.title = "README"
        banner(ws, "Prairie Home Furnishings (fictional) - Integrated Three-Statement Financial Model", "Demonstration developed by VisiVault Analytics Ltd. | October 2026", 8)
        lines = [
            "WHAT THIS IS",
            "SIMULATED accounts of a FICTIONAL Calgary furniture retailer, built by double entry from synthetic operational data. Not actual company records.",
            "Historical months (Jan-2024 to Sep-2026) are simulated actuals from the general ledger; Oct-2026 to Sep-2027 are FORECASTS.",
            "",
            "HOW TO USE",
            "1. Assumptions!C4 - choose Conservative / Base / Optimistic / Custom (edit the Custom column freely).",
            "2. Drivers - edit any blue monthly input (bookings, backlog release, staffing plan, lease, capex plan, inventory cover policy).",
            "3. Forecast - three statements recalculate; the CHECKS rows and the Checks sheet must stay at 0 / TRUE.",
            "4. Scenario_Results - stored comparison of the three scenarios and the sensitivity cases.",
            "",
            "ACCOUNTING POLICIES",
            "Revenue is recognised at hand-off (store pickup or loading to the customer's carrier), not at order. Deposits and balances received earlier are customer deposits",
            "(contract liabilities). Third-party delivery is paid by the customer to the carrier and is not revenue; flat-rate delivery and protection plans are collected as agent",
            "(only the plan commission is revenue). Inventory at purchase cost; stocked goods and goods allocated to customer orders are shown separately.",
            "",
            "CONVENTIONS",
            "CAD. Income-statement and cash-flow lines are signed: income/inflows positive, costs/outflows in (parentheses). Blue-on-yellow cells are inputs; black cells are formulas.",
            "Interest is charged on opening balances and the revolver follows a cash-band rule, so the model has no circular references and no balancing plugs.",
            "",
            "SHEETS",
            "Assumptions | Drivers | Forecast | Scenario_Results | Checks | Hist_IS | Hist_BS | Hist_CF | Hist_Balances",
        ]
        for i, l in enumerate(lines):
            c = ws.cell(row=5 + i, column=1, value=l)
            if l.isupper():
                c.font = Font(bold=True, color=BLUE)
        ws.column_dimensions["A"].width = 160

    def build(self, path: Path):
        self.historicals()
        self.assumptions()
        self.drivers()
        self.forecast_sheet()
        self.scenario_sheet()
        self.checks_sheet()
        self.readme()
        order = ["README", "Assumptions", "Drivers", "Forecast", "Scenario_Results", "Checks", "Hist_IS", "Hist_BS", "Hist_CF", "Hist_Balances"]
        self.wb._sheets = [self.wb[n] for n in order]
        self.wb.save(path)
        return path


# ======================================================================================== value workbooks
def _sheet(wb, name, title, df, sub=None, subtotal_rows=(), fmt=FMT, index=True):
    ws = wb.create_sheet(name[:31])
    banner(ws, title, sub, max(8, len(df.columns) + 1))
    write_df(ws, df, 5, index=index, fmt=fmt, subtotal_rows=subtotal_rows)
    return ws


def labelled(df, labels):
    d = df.copy()
    d.index = [labels.get(k, k) for k in d.index]
    d.index.name = "Line (CAD)"
    return d


def statements_workbook(L_, S_, path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "README"
    banner(ws, "Prairie Home Furnishings (fictional) - Monthly, quarterly and annual financial statements", "SIMULATED - generated by double entry from synthetic operational data.")
    ws["A5"] = "Sheets: IS/BS/CF monthly, quarterly and annual; trial balances; chart of accounts; subsidiary-ledger reconciliation; AR/AP ageing; fixed assets; debt; inventory valuation."
    ws["A6"] = "CAD. Income-statement and cash-flow lines signed (income/inflow +, cost/outflow -). Year-end balance sheets are post-closing."
    subs_is = {S.IS_LABEL[k] for k in IS_SUB}
    subs_bs = {S.BS_LABEL[k] for k in BS_SUB} | {S.BS_LABEL["check"]}
    subs_cf = {S.CF_LABEL[k] for k in list(CF_SUB) + ["close_cash", "check"]}
    for per, tag in (("m", "Monthly"), ("q", "Quarterly"), ("y", "Annual")):
        _sheet(wb, f"IS_{tag}", f"Income statement - {tag.lower()} (CAD)", labelled(S_[f"is_{per}"], S.IS_LABEL), subtotal_rows=subs_is)
        _sheet(wb, f"BS_{tag}", f"Balance sheet - {tag.lower()} period end (CAD)", labelled(S_[f"bs_{per}"], S.BS_LABEL), subtotal_rows=subs_bs)
        _sheet(wb, f"CF_{tag}", f"Cash flow statement (indirect) - {tag.lower()} (CAD)", labelled(S_[f"cf_{per}"], S.CF_LABEL), subtotal_rows=subs_cf)
    _sheet(wb, "Opening_BS_31Dec2023", "Opening statement of financial position 31-Dec-2023 (retained earnings = derived residual)", labelled(S_["bs_open"].to_frame("31-Dec-2023"), S.BS_LABEL))
    _sheet(wb, "Direct_CF_crosscheck", "Direct-method cash flows (cash-account lines by journal source) - cross-check", S_["direct"])
    tb = S_["tb"]
    piv = tb.assign(net=tb.debit - tb.credit).pivot_table(index=["account", "account_name"], columns="month", values="net", aggfunc="sum", fill_value=0.0)
    piv.loc[("", "TOTAL (must be 0)"), :] = piv.sum().round(2)
    piv.index = [f"{a} {n}".strip() for a, n in piv.index]
    _sheet(wb, "Trial_Balances", "Monthly trial balances, pre-closing (debit +, credit -)", piv, fmt=FMT2, subtotal_rows={"TOTAL (must be 0)"})
    coa = pd.DataFrame(P.COA, columns=["code", "name", "type", "normal_balance", "statement_line", "cash_flow_class"])
    _sheet(wb, "Chart_of_Accounts", "Chart of accounts", coa, index=False)
    _sheet(wb, "Subledger_Reconciliation", "AR and AP subsidiary ledgers vs control accounts (difference must be 0)", S_["subledger"], index=False, fmt=FMT2)
    for nm, items in (("AR", L_["ar_items"]), ("AP", L_["ap_items"])):
        ag = S.aging(items)
        if len(ag):
            t = ag.pivot_table(index="counterparty", columns="bucket", values="amount", aggfunc="sum", fill_value=0.0, observed=False)
            t["Total"] = t.sum(axis=1)
            _sheet(wb, f"{nm}_Ageing_30Sep2026", f"{nm} open items at 30-Sep-2026 by age (days since document)", t.sort_values("Total", ascending=False), fmt=FMT2)
    _sheet(wb, "Fixed_Asset_Register", "Fixed-asset register and depreciation (straight line)", S_["far"], index=False, fmt=FMT2)
    _sheet(wb, "Debt_Treasury_Schedule", "Debt, interest, revolver, tax and treasury schedule (month end)", L_["treasury"], index=False, fmt=FMT2)
    iv = S_["inv_val"].copy()
    iv["gl_stock_1200"] = S_["bs_m"].loc["inv_stock"].values
    iv["gl_allocated_1210"] = S_["bs_m"].loc["inv_alloc"].values
    iv["difference"] = (iv.gl_stock_1200 - iv.stock_value_from_movements + iv.gl_allocated_1210 - iv.allocated_value_from_movements).round(2)
    _sheet(wb, "Inventory_Valuation", "Inventory valuation: movement records vs general ledger (stocked and customer-allocated)", iv, fmt=FMT2)
    _sheet(wb, "Common_Size_Annual", "Common-size income statement (% of total net revenue)", labelled(S_["cs_y"], S.IS_LABEL), fmt=FMT2)
    wb.save(path)
    return path


def ratios_workbook(rat: pd.DataFrame, ann: pd.DataFrame, path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Annual_Comparison"
    banner(ws, "Financial ratio analysis - FY2024, FY2025 and TTM Sep-2026 (SIMULATED)", "Ratios computed from the illustrative management accounts; see Definitions.")
    a = ann.copy()
    a["management_action"] = a.index.map(R.ACTIONS)
    write_df(ws, a.reset_index(drop=True), 5, index=False, fmt=FMT2)
    lab = {d[0]: f"{d[1]} ({d[5]})" for d in R.DEFINITIONS}
    _sheet(wb, "Monthly_TTM", "Monthly ratio series (TTM flows; month-end / 13-point average balances)", labelled(rat, lab), fmt=FMT2)
    defs = pd.DataFrame(R.DEFINITIONS, columns=["key", "ratio", "group", "formula", "meaning", "unit", "higher_is_better"])
    _sheet(wb, "Definitions", "Ratio definitions", defs, index=False)
    wb.save(path)
    return path


def working_capital_workbook(S_, rat, FC, path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Monthly_Working_Capital"
    banner(ws, "Working capital, inventory and customer-deposit analysis (SIMULATED)", "Month-end balances from the general ledger.")
    bs = S_["bs_m"]
    wc = pd.DataFrame({"Receivables": bs.loc["ar"], "Inventory - stocked": bs.loc["inv_stock"], "Inventory - allocated to customers": bs.loc["inv_alloc"],
                       "Prepaid": bs.loc["prepaid"], "Accounts payable": -bs.loc["ap"], "Accrued": -bs.loc["accrued"], "Customer deposits": -bs.loc["deposits"],
                       "Agent liabilities": -bs.loc["agent"], "Net sales tax": bs.loc["gst_rec"] - bs.loc["gst_pay"]})
    wc["Net operating working capital"] = wc.sum(axis=1)
    wc["Cash"] = bs.loc["cash"]
    wc["Deposits / total inventory"] = bs.loc["deposits"] / (bs.loc["inv_stock"] + bs.loc["inv_alloc"])
    wc.index.name = "Month"
    write_df(ws, wc, 5, fmt=FMT)
    cyc = rat.loc[["dio", "dso", "dpo", "ccc", "ccc_after_deposits", "inventory_turnover", "current_ratio", "quick_ratio", "deposit_funding"]].T
    cyc.index.name = "Month (TTM)"
    _sheet(wb, "Cash_Conversion_Cycle", "Cash conversion cycle (days) before and after customer deposits, TTM basis", cyc, fmt=FMT2)
    f = FC["results"]["Base"]
    fwc = f.loc[["bookings", "net_sales", "backlog", "deposits", "inv_stock", "inv_alloc", "purchases", "ap", "working_capital", "cash"]].copy()
    fwc.index = ["Bookings", "Net sales recognised", "Open order backlog", "Customer deposits", "Inventory - stocked", "Inventory - allocated", "Inventory purchases",
                 "Accounts payable", "Working capital", "Cash"]
    _sheet(wb, "Forecast_Working_Capital", "Forecast working capital and cash - Base scenario (FORECAST)", fwc)
    wb.save(path)
    return path


def forecast_workbook(FC, path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scenario_Summary"
    banner(ws, "12-month financial forecast Oct-2026 to Sep-2027 - scenario values (FORECAST, SIMULATED)",
           "For live recalculation use the Integrated Financial Model workbook.")
    write_df(ws, FC["summary"].drop(columns="assumptions").T, 5)
    labels = {k: lab for k, lab, *_ in F.ROWS}
    for name, res in FC["results"].items():
        _sheet(wb, f"{name}_Monthly", f"{name} scenario - monthly integrated forecast (CAD)", labelled(res, labels))
    _sheet(wb, "Sensitivity", "Sensitivity analysis vs Base (one driver at a time)", FC["sensitivity"], index=False)
    notes = pd.DataFrame({"scenario": list(F.SCENARIO_NOTES), "assumptions": list(F.SCENARIO_NOTES.values())})
    _sheet(wb, "Scenario_Assumptions", "Scenario definitions", notes, index=False)
    wb.save(path)
    return path


# ======================================================================================== Excel recalculation (verification)
APPLESCRIPT = """on run argv
  set f to POSIX file (item 1 of argv)
  tell application "Microsoft Excel"
    set wb to open workbook workbook file name f
    calculate full
    save wb
    close wb saving no
  end tell
end run"""


def excel_available() -> bool:
    return Path("/Applications/Microsoft Excel.app").exists() and shutil.which("osascript") is not None


def recalc_with_excel(path: Path, scratch: Path) -> bool:
    """Open, fully recalculate and save a workbook in Microsoft Excel (macOS) so cached values can be read back."""
    scr = scratch / "recalc.applescript"
    scr.write_text(APPLESCRIPT)
    r = subprocess.run(["osascript", str(scr), str(path.resolve())], capture_output=True, text=True, timeout=600)
    return r.returncode == 0


def set_scenario(src: Path, dst: Path, name: str, edits: dict | None = None):
    wb = openpyxl.load_workbook(src)
    wb["Assumptions"]["C4"] = name
    for (sheet, cell), v in (edits or {}).items():
        wb[sheet][cell] = v
    wb.save(dst)
